import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

try:
    from openpyxl import load_workbook
    import serial
except ImportError:
    load_workbook = None
    serial = None
from tools.record_sensor_test import capture, main, parse_line, save_report, summarize


@unittest.skipUnless(load_workbook is not None and serial is not None,
                     'Install tools/requirements-sensor-test.txt for recorder tests')
class SensorRecorderTests(unittest.TestCase):
    def test_both_sketch_formats_and_failures(self):
        self.assertEqual(parse_line('Distance: 25.4 cm'), ('READING', 25.4))
        self.assertEqual(parse_line('Distance: 25.4 cm | echo: 1481 us'), ('READING', 25.4))
        self.assertEqual(parse_line('Distance: 501.0 cm | OUT OF STATED SENSOR RANGE'), ('OUT OF RANGE', 501.0))
        self.assertEqual(parse_line('ECHO already HIGH: check wiring and divider.'), ('ECHO HIGH', None))
        self.assertEqual(parse_line('OUT OF RANGE: outside 2-200 cm; reading rejected.'), ('OUT OF RANGE', None))
        self.assertEqual(parse_line('Distance: broken'), ('MALFORMED', None))
        self.assertIsNone(parse_line('Accepted range: 2-200 cm. Other distances report OUT OF RANGE.'))

    def test_fragmented_serial_counts_errors_and_preserves_bytes(self):
        chunks = [b'HC-SR04 sensor test\r\nDista', b'nce: 24.0 cm\r\nNO E',
                  b'CHO: check target\nOUT OF RANGE: rejected\nDistance: 26.0 cm\n']
        stream = Mock(); stream.read.side_effect = chunks
        log = io.BytesIO(); records = []
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(capture(stream, log, records, 4, 1), 'complete')
        self.assertEqual(log.getvalue(), b''.join(chunks))
        summary = summarize(records, 25)
        self.assertEqual(summary['recorded_attempts'], 4)
        self.assertEqual(summary['return_rate'], .5)
        self.assertEqual(summary['mean_cm'], 25)
        self.assertEqual(summary['bias_cm'], 0)
        self.assertEqual(summary['mean_absolute_error_cm'], 1)
        self.assertAlmostEqual(summary['sample_sd_cm'], 2**.5)
        self.assertEqual(summary['out_of_range'], 1)

    def test_empty_and_all_failure_statistics_are_not_zero_distances(self):
        for records in [[], [{'status': 'NO ECHO', 'distance_cm': None}]]:
            result = summarize(records, 25)
            self.assertIsNone(result['mean_cm'])
            self.assertIsNone(result['sample_sd_cm'])
            self.assertIsNone(result['bias_cm'])

    def test_timeout_preserves_partial_sample(self):
        stream = Mock(); stream.read.return_value = b'Distance: 24.0 cm\n'
        records = []
        with patch('tools.record_sensor_test.time.monotonic', side_effect=[0, 0, 2]):
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(capture(stream, io.BytesIO(), records, 30, 1), 'timeout')
        self.assertEqual(len(records), 1)

    def test_excel_contains_real_rows_statistics_and_literal_notes(self):
        records = [dict(sample=1, received_at='2026-09-14T12:00:00+10:00', status='READING', distance_cm=24.5, raw_line='Distance: 24.5 cm'),
                   dict(sample=2, received_at='2026-09-14T12:00:01+10:00', status='NO ECHO', distance_cm=None, raw_line='NO ECHO: check target')]
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            save_report(folder, records, {'sensor': 'A', 'true_distance_cm': 25, 'notes': '=1+1\x00', 'capture_status': 'timeout'})
            book = load_workbook(folder / 'results.xlsx')
            fields = {row[0].value: row[1] for row in book['Summary'].iter_rows(min_row=2)}
            self.assertEqual(fields['mean_cm'].value, 24.5)
            self.assertEqual(fields['bias_cm'].value, -.5)
            self.assertEqual(fields['return_rate'].value, .5)
            self.assertEqual(fields['notes'].data_type, 's')
            self.assertEqual(fields['notes'].value, '=1+1')
            self.assertEqual(book['Readings'].max_row, 3)
            self.assertIsNone(book['Readings']['D3'].value)
            self.assertEqual(book['Readings']['C3'].value, 'NO ECHO')
            self.assertTrue((folder / 'readings.csv').exists())
            self.assertEqual(json.loads((folder / 'run.json').read_text())['capture_status'], 'timeout')
            book.close()

    def test_disconnect_saves_partial_evidence(self):
        stream = Mock()
        stream.read.side_effect = [b'Distance: 25.0 cm\n', serial.SerialException('disconnected')]
        with tempfile.TemporaryDirectory() as tmp:
            with patch('serial.Serial', return_value=stream), patch('serial.tools.list_ports.comports', return_value=[]), patch('tools.record_sensor_test.time.sleep'):
                with contextlib.redirect_stdout(io.StringIO()):
                    code = main(['--distance-cm', '25', '--port', 'test-port', '--output', tmp])
            self.assertEqual(code, 1)
            folder, = Path(tmp).iterdir()
            meta = json.loads((folder / 'run.json').read_text())
            self.assertEqual(meta['capture_status'], 'serial_error')
            self.assertEqual(meta['summary']['recorded_attempts'], 1)
            self.assertEqual((folder / 'serial.log').read_bytes(), b'Distance: 25.0 cm\n')
            self.assertTrue((folder / 'results.xlsx').exists())
        stream.close.assert_called_once()

    def test_no_port_creates_no_test_results(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch('serial.tools.list_ports.comports', return_value=[]), contextlib.redirect_stderr(io.StringIO()):
                with self.assertRaises(SystemExit) as exc:
                    main(['--distance-cm', '25', '--output', tmp])
            self.assertEqual(exc.exception.code, 2)
            self.assertEqual(list(Path(tmp).iterdir()), [])


if __name__ == '__main__':
    unittest.main()
