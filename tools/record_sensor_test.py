"""Capture the Arduino sensor_test USB output into local Excel/CSV evidence."""
import argparse
from collections import Counter
import csv
from datetime import datetime
import hashlib
import json
import math
import os
from pathlib import Path
import re
import statistics
import sys
import time
import uuid

ROOT = Path(__file__).resolve().parents[1]
DISTANCE = re.compile(r"^Distance:\s*([0-9]+(?:\.[0-9]+)?)\s+cm(?:\s*\|.*)?$")


def parse_line(line):
    """Return (status, distance) for one attempt; ignore startup chatter."""
    line = line.strip()
    match = DISTANCE.fullmatch(line)
    if match:
        value = float(match.group(1))
        if not math.isfinite(value):
            return "MALFORMED", None
        # Retain reported numeric values even if the older sketch flags them.
        status = "OUT OF RANGE" if "OUT OF" in line else "READING"
        return status, value
    for prefix, status in [("NO ECHO", "NO ECHO"),
                           ("ECHO already HIGH", "ECHO HIGH"),
                           ("OUT OF RANGE", "OUT OF RANGE")]:
        if line.startswith(prefix):
            return status, None
    if line.startswith("Distance:"):
        return "MALFORMED", None
    return None


def capture(stream, raw_log, records, samples, timeout):
    """Keep every error attempt; preserve serial bytes and split-line fragments."""
    deadline = time.monotonic() + timeout
    pending = b""
    while len(records) < samples and time.monotonic() < deadline:
        chunk = stream.read(256)
        if not chunk:
            continue
        raw_log.write(chunk)
        raw_log.flush()
        pending += chunk
        while b"\n" in pending and len(records) < samples:
            line, pending = pending.split(b"\n", 1)
            line = line.decode("utf-8", errors="replace").strip()
            parsed = parse_line(line)
            if parsed is None:
                continue
            status, distance = parsed
            records.append({"sample": len(records) + 1,
                            "received_at": datetime.now().astimezone().isoformat(),
                            "status": status, "distance_cm": distance, "raw_line": line})
            print(f"{len(records):>3}/{samples}: {line}", flush=True)
        # A corrupt stream must not grow memory without bound. Bytes stay in raw log.
        if len(pending) > 65536:
            pending = b""
    return "complete" if len(records) == samples else "timeout"


def summarize(records, true_cm):
    counts = Counter(r["status"] for r in records)
    values = [r["distance_cm"] for r in records if r["status"] == "READING"]
    return {
        "recorded_attempts": len(records), "numerical_returns": len(values),
        "no_echo": counts["NO ECHO"], "echo_high": counts["ECHO HIGH"],
        "out_of_range": counts["OUT OF RANGE"], "malformed": counts["MALFORMED"],
        "return_rate": len(values) / len(records) if records else None,
        "mean_cm": statistics.mean(values) if values else None,
        "bias_cm": statistics.mean(values) - true_cm if values else None,
        "sample_sd_cm": statistics.stdev(values) if len(values) > 1 else None,
        "min_cm": min(values) if values else None, "max_cm": max(values) if values else None,
        "mean_absolute_error_cm": statistics.mean(abs(v - true_cm) for v in values) if values else None,
    }


def save_report(folder, records, metadata):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.chart import LineChart, Reference
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def excel_value(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    summary = summarize(records, metadata["true_distance_cm"])
    (folder / "run.json").write_text(json.dumps({**metadata, "summary": summary}, indent=2), encoding="utf-8")
    headers = ["Sample", "Received at", "Status", "Reading (cm)", "True distance (cm)",
               "Bias (cm)", "Absolute error (cm)", "Original serial line"]
    rows = []
    for r in records:
        error = r["distance_cm"] - metadata["true_distance_cm"] if r["status"] == "READING" else None
        rows.append([r["sample"], r["received_at"], r["status"], r["distance_cm"],
                     metadata["true_distance_cm"], error, abs(error) if error is not None else None, r["raw_line"]])
    with (folder / "readings.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle); writer.writerow(headers); writer.writerows(rows)
    book = Workbook()
    result = book.active; result.title = "Summary"
    result.append(["Measure", "Result"])
    for key, value in {**metadata, **summary}.items():
        result.append([key, excel_value(value)])
    result.append(["Interpretation", "Capture completion is not a reliability pass. Numerical returns may still be inaccurate."])
    result.append(["Statistics", "Statistics use READING rows only. Error statuses remain in the attempt counts and raw evidence."])
    result.append(["Range filtering", "This recorder does not impose a distance cap. OUT OF RANGE comes from the uploaded sketch."])
    result.append(["Calibration", "Bias is measured minus true distance. No correction is applied to the readings."])
    result.append(["Raw evidence", "serial.log preserves bytes received during capture; the final chunk can include extra uncounted lines."])
    data = book.create_sheet("Readings"); data.append(headers)
    for row in rows:
        data.append([excel_value(value) for value in row])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for row in sheet.iter_rows():
            for cell in row:
                # Keep serial text/notes as literal evidence, never Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    result.column_dimensions["A"].width = 29; result.column_dimensions["B"].width = 110
    for col, width in zip("ABCDEFGH", [12, 36, 20, 19, 22, 17, 23, 85]):
        data.column_dimensions[col].width = width
    for row in data.iter_rows(min_row=2):
        for cell in row[3:7]:
            cell.number_format = "0.00"
    if records:
        chart = LineChart(); chart.title = "Reported distance and known target distance"
        chart.y_axis.title = "Distance (cm)"; chart.x_axis.title = "Attempt"
        chart.add_data(Reference(data, min_col=4, max_col=5, min_row=1, max_row=len(rows)+1), titles_from_data=True)
        chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=len(rows)+1))
        chart.display_blanks = "gap"
        result.add_chart(chart, "D2")
    book.save(folder / "results.xlsx")
    return summary


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sensor", choices=["A", "B"], default="A")
    parser.add_argument("--distance-cm", type=float, help="Measured target gap from transducer faces")
    parser.add_argument("--samples", type=int, default=30)
    parser.add_argument("--port", help="Serial port; auto-selects only when exactly one USB serial device is found")
    parser.add_argument("--list-ports", action="store_true")
    parser.add_argument("--timeout", type=float, default=60, help="Capture deadline in seconds; saves incomplete runs")
    parser.add_argument("--notes", default="", help="Target, power supply, temperature, etc.")
    parser.add_argument("--output", type=Path, default=ROOT / "evidence")
    args = parser.parse_args(argv)
    try:
        import serial
        from serial.tools import list_ports
        import openpyxl  # noqa: F401 - check dependency before capture begins
    except ImportError:
        parser.exit(2, "Install recorder dependencies: python3 -m pip install -r tools/requirements-sensor-test.txt\n")
    ports = list(list_ports.comports())
    if args.list_ports:
        for port in ports:
            print(f"{port.device}: {port.description}")
        if not ports:
            print("No serial ports found. Connect the ESP32 with a USB data cable.")
        return 0
    if args.distance_cm is None or not math.isfinite(args.distance_cm) or not 2 <= args.distance_cm <= 500:
        parser.error("--distance-cm must be a measured distance from 2 to 500 cm")
    if not 1 <= args.samples <= 10000 or not math.isfinite(args.timeout) or not 1 <= args.timeout <= 3600:
        parser.error("Use 1..10000 samples and a 1..3600 second timeout")
    if not args.port:
        candidates = [p for p in ports if p.vid is not None or re.search(r"tty(?:USB|ACM)|cu\.usb", p.device)]
        if len(candidates) != 1:
            parser.exit(2, "Connect one ESP32, or use --list-ports and specify --port. No test was recorded.\n")
        args.port = candidates[0].device
    print(f"Sensor {args.sensor}, target {args.distance_cm:g} cm, {args.samples} attempts on {args.port}.")
    print("Close Arduino Serial Monitor/Plotter. Keep the target still and the servo unpowered.", flush=True)
    try:
        stream = serial.Serial(port=None, baudrate=115200, timeout=.25,
                               exclusive=True if os.name == "posix" else None)
        stream.port = args.port
        stream.dtr = False; stream.rts = False
        stream.open()
    except (serial.SerialException, OSError) as exc:
        parser.exit(2, f"Cannot open serial port: {exc}\nClose Serial Monitor and check Tools > Port.\n")
    records = []
    try:
        # Opening some USB adapters resets the board. Start the sample window after settling.
        time.sleep(2)
        stream.reset_input_buffer()
        started = datetime.now().astimezone()
        folder = args.output.resolve() / f"{started:%Y%m%d_%H%M%S}_sensor-{args.sensor}_{args.distance_cm:g}cm_{uuid.uuid4().hex[:8]}"
        folder.mkdir(parents=True, exist_ok=False)
        metadata = {"sensor": args.sensor, "true_distance_cm": args.distance_cm,
                    "requested_attempts": args.samples, "started_at": started.isoformat(),
                    "port": args.port, "baud": 115200, "capture_timeout_s": args.timeout,
                    "startup_excluded_s": 2, "notes": args.notes,
                    "recorder_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                    "uploaded_firmware": "Not verified over serial; use Arduino sensor_test sketch."}
        print(f"Recording now. Evidence: {folder}", flush=True)
        with (folder / "serial.log").open("wb") as raw_log:
            try:
                outcome = capture(stream, raw_log, records, args.samples, args.timeout)
            except KeyboardInterrupt:
                outcome = "interrupted"
            except (serial.SerialException, OSError) as exc:
                outcome = "serial_error"
                metadata["serial_error"] = str(exc)
        metadata.update(finished_at=datetime.now().astimezone().isoformat(), capture_status=outcome)
        summary = save_report(folder, records, metadata)
    finally:
        stream.close()
    print(f"Saved Excel: {folder / 'results.xlsx'}")
    print(f"Capture: {outcome}; attempts: {len(records)}/{args.samples}; numerical returns: {summary['numerical_returns']}.")
    if summary["mean_cm"] is not None:
        print(f"Mean: {summary['mean_cm']:.2f} cm; bias: {summary['bias_cm']:+.2f} cm; MAE: {summary['mean_absolute_error_cm']:.2f} cm.")
    return 0 if outcome == "complete" else 1


if __name__ == "__main__":
    sys.exit(main())
